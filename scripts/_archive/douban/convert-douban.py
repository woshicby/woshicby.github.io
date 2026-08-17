#!/usr/bin/env python3
"""豆瓣豆伴 Excel → JSON 转换器(Python 版,替代原 convert-douban.js)。

用法:
    ./run.sh convert-douban.py          # 或: .venv/bin/python convert-douban.py

读取 documents/豆伴(232919949).xlsx,按 sheet 顺序映射状态,
按链接识别分类(movie/music/book/game/drama),解析 intro,
输出 JSON/review-{category}s.json。

依赖: openpyxl(已在 scripts/requirements.txt)
"""
import json
import os
import re
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)

# sheet 顺序 → 状态(sheet 索引即豆瓣导出顺序)
SHEET_STATUS_MAP = [
    {"status": "watched"}, {"status": "watching"}, {"status": "wantToWatch"},
    {"status": "listened"}, {"status": "listening"}, {"status": "wantToListen"},
    {"status": "read"}, {"status": "reading"}, {"status": "wantToRead"},
    {"status": "played"}, {"status": "playing"}, {"status": "wantToPlay"},
    {"status": "watched"}, {"status": "wantToWatch"},
]

KNOWN_PLATFORMS = [
    "PC", "Mac", "iPhone", "iPad", "Android", "PlayStation 5", "PlayStation 4",
    "Xbox Series X", "Xbox One", "Nintendo Switch", "Nintendo Switch 2", "Steam Deck",
    "Web", "PS Vita", "Nintendo 3DS", "Wii U",
]


def detect_category(link):
    if not link:
        return "unknown"
    if "movie.douban.com" in link:
        return "movie"
    if "music.douban.com" in link:
        return "music"
    if "book.douban.com" in link:
        return "book"
    if "douban.com/game" in link:
        return "game"
    if "douban.com/location/drama" in link:
        return "drama"
    return "unknown"


def parse_movie_intro(intro):
    if not intro:
        return {"year": None, "region": "", "genres": [], "directors": "", "actors": ""}
    parts = intro.split(" / ")
    return {
        "year": int(parts[0]) if parts[0].isdigit() else None,
        "region": parts[1] if len(parts) > 1 else "",
        "genres": parts[2].split(" ") if len(parts) > 2 and parts[2] else [],
        "directors": parts[3] if len(parts) > 3 else "",
        "actors": parts[4] if len(parts) > 4 else "",
    }


def parse_book_intro(intro):
    if not intro:
        return {"author": "", "year": None, "publisher": ""}
    parts = intro.split(" / ")
    return {
        "author": parts[0] if parts else "",
        "year": int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else None,
        "publisher": parts[2] if len(parts) > 2 else "",
    }


def parse_music_intro(intro):
    if not intro:
        return {"artist": "", "year": None}
    parts = intro.split(" / ")
    return {
        "artist": parts[0] if parts else "",
        "year": int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else None,
    }


def parse_game_intro(intro):
    if not intro:
        return {"genres": [], "platforms": [], "developer": "", "releaseDate": ""}
    parts = intro.split("/")
    genres, platforms = [], []
    developer, release_date = "", ""

    for part in parts:
        trimmed = part.strip()
        if not trimmed:
            continue
        if trimmed in KNOWN_PLATFORMS:
            platforms.append(trimmed)
        elif re.fullmatch(r"\d{4}-\d{2}-\d{2}", trimmed):
            release_date = trimmed
        elif trimmed == trimmed.upper() and len(trimmed) > 2 and " " not in trimmed:
            developer = trimmed
        elif re.match(r"^[A-Z]", trimmed) and " " in trimmed and trimmed not in KNOWN_PLATFORMS:
            developer = trimmed
        elif re.match(r"^[\u4e00-\u9fa5]", trimmed) and trimmed not in KNOWN_PLATFORMS:
            if len(genres) == 0 and not developer:
                genres.append(trimmed)
            elif not developer and not release_date:
                developer = trimmed
            else:
                genres.append(trimmed)
        else:
            genres.append(trimmed)

    return {"genres": genres, "platforms": platforms, "developer": developer, "releaseDate": release_date}


def parse_drama_intro(intro):
    if not intro:
        return {"type": "", "name": ""}
    parts = intro.split(" / ")
    return {
        "type": parts[0].strip() if parts else "",
        "name": parts[1].strip() if len(parts) > 1 else "",
    }


def map_row(row, category, status):
    keys = list(row.keys())
    def get(idx):
        return str(row[keys[idx]]) if idx < len(keys) and row[keys[idx]] is not None else ""

    title = get(0)
    intro = get(1)
    douban_rating = float(get(2)) if get(2).replace(".", "", 1).isdigit() else 0
    link = get(3)
    created_at = get(4)
    my_rating = int(get(5)) if get(5).isdigit() else 0
    tags_str = get(6)
    review = get(7)

    tags = [t for t in tags_str.split(" ") if t] if tags_str else []
    base = {"title": title, "doubanRating": douban_rating, "myRating": my_rating,
            "review": review, "tags": tags, "link": link, "createdAt": created_at,
            "category": category, "status": status}

    if category == "movie":
        p = parse_movie_intro(intro)
        base.update({"year": p["year"], "region": p["region"], "genres": p["genres"],
                     "directors": p["directors"], "actors": p["actors"]})
    elif category == "book":
        p = parse_book_intro(intro)
        base.update({"author": p["author"], "year": p["year"], "publisher": p["publisher"]})
    elif category == "music":
        p = parse_music_intro(intro)
        base.update({"artist": p["artist"], "year": p["year"]})
    elif category == "game":
        p = parse_game_intro(intro)
        base.update({"genres": p["genres"], "platforms": p["platforms"],
                     "developer": p["developer"], "releaseDate": p["releaseDate"]})
    elif category == "drama":
        p = parse_drama_intro(intro)
        base.update({"type": p["type"], "dramaName": p["name"]})
    return base


def convert():
    import openpyxl

    xlsx_path = os.path.join(PROJECT_ROOT, "documents", "豆伴(232919949).xlsx")
    if not os.path.exists(xlsx_path):
        print(f"❌ 找不到豆伴 Excel: {xlsx_path}")
        sys.exit(1)

    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    result = {}

    for sheet_index, sheet_name in enumerate(workbook.sheetnames):
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        header = rows[0]
        data_rows = rows[1:]
        if not data_rows:
            continue

        status_info = SHEET_STATUS_MAP[sheet_index] if sheet_index < len(SHEET_STATUS_MAP) else {"status": "unknown"}
        status = status_info["status"]

        # 从第一行第 4 列(link)判断分类
        first_link = str(data_rows[0][3]) if len(data_rows[0]) > 3 and data_rows[0][3] else ""
        category = detect_category(first_link)

        result.setdefault(category, {})
        items = []
        for row in data_rows:
            # 转成 dict(用列名或列索引作 key,保持与 JS 版 getValueByIndex 逻辑一致)
            row_dict = {i: (row[i] if i < len(row) else "") for i in range(len(header))}
            items.append(map_row(row_dict, category, status))
        result[category][status] = items
        print(f"Sheet {sheet_index}: -> {category}/{status} ({len(items)} items)")

    for category, statuses in result.items():
        output_path = os.path.join(PROJECT_ROOT, "JSON", f"review-{category}s.json")
        with open(output_path, "w", encoding="utf-8-sig") as f:
            json.dump(statuses, f, ensure_ascii=False, indent=2)
        total = sum(len(arr) for arr in statuses.values())
        print(f"Written: review-{category}s.json ({total} total items)")

    print("\nDone!")


if __name__ == "__main__":
    convert()
